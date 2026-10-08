import json
import sys
from pathlib import Path
from uuid import uuid4

# Ensure backend root is on sys.path for direct execution
BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from sqlalchemy import text

from app.db.session import engine
from app.db.queries import (
    check_db_connection,
    get_all_table_names,
    get_table_data,
    get_all_database_data,
)
from app.db.export_excel import export_database_to_excel


def test_database_connection_contact() -> None:
    """Test establishing contact with the database via raw SQL query."""
    info = check_db_connection()

    assert info is not None
    assert "db_name" in info
    assert info["db_name"] == "postgres"
    assert "user_name" in info
    assert "pg_version" in info


def test_sql_query_get_all_table_names() -> None:
    """Test retrieving all table names in public schema using SQL query."""
    tables = get_all_table_names()

    assert isinstance(tables, list)
    # Core registered application and migration tables must exist
    expected_tables = {
        "alembic_version",
        "quiz_results",
        "game_results",
        "tutor_strategy_predictions",
        "game_sessions",
        "gameplay_events",
        "game_task_attempts",
    }
    assert expected_tables.issubset(set(tables))


def test_sql_query_get_all_db_data() -> None:
    """Test executing SQL queries to fetch all data across all tables."""
    all_data = get_all_database_data()

    assert isinstance(all_data, dict)
    assert "alembic_version" in all_data
    # alembic_version must contain the current applied migration record
    assert len(all_data["alembic_version"]) >= 1
    assert "version_num" in all_data["alembic_version"][0]

    # Verify every registered table is represented as a list of record mappings
    for table_name in [
        "quiz_results",
        "game_results",
        "tutor_strategy_predictions",
        "game_sessions",
        "gameplay_events",
        "game_task_attempts",
    ]:
        assert table_name in all_data
        assert isinstance(all_data[table_name], list)


def test_sql_query_individual_table() -> None:
    """Test querying data from a specific table with raw SQL."""
    alembic_rows = get_table_data("alembic_version")
    assert isinstance(alembic_rows, list)
    assert len(alembic_rows) > 0
    assert "version_num" in alembic_rows[0]


def test_transactional_insert_and_sql_query_roundtrip() -> None:
    """Test contact with DB by inserting sample data and querying it with SQL."""
    session_id = f"test_session_{uuid4().hex[:8]}"
    test_id = str(uuid4())

    with engine.connect() as conn:
        with conn.begin():
            # 1. Insert test record using SQL
            insert_sql = text(
                """
                INSERT INTO quiz_results (
                    id, quiz_session_id, student_id, concept_id,
                    weak_concept, quiz_mastery_score, recommended_support_level,
                    questions_total, questions_attempted, correct_answers,
                    quiz_accuracy, quiz_avg_response_time_sec, quiz_hint_rate,
                    misconception_code, quiz_difficulty_level, quiz_attempt_count,
                    completed_at
                ) VALUES (
                    :id, :session_id, 'student_test', 'math_add_01',
                    'math_add_01', 0.85, 'medium',
                    5, 5, 5,
                    1.0, 12.5, 0.0,
                    'none', 'intermediate', 1,
                    NOW()
                )
                """
            )
            conn.execute(insert_sql, {"id": test_id, "session_id": session_id})

            # 2. Query all data from quiz_results with SQL
            select_sql = text("SELECT * FROM quiz_results WHERE quiz_session_id = :session_id")
            fetched = conn.execute(select_sql, {"session_id": session_id}).mappings().all()

            assert len(fetched) == 1
            record = dict(fetched[0])
            assert record["quiz_session_id"] == session_id
            assert record["student_id"] == "student_test"
            assert record["concept_id"] == "math_add_01"
            assert float(record["quiz_accuracy"]) == 1.0

            # 3. Roll back transaction so test leaves no side-effects
            conn.rollback()

    # 4. Verify after rollback that the test record does not persist
    with engine.connect() as conn:
        check_sql = text("SELECT COUNT(*) FROM quiz_results WHERE quiz_session_id = :session_id")
        count = conn.execute(check_sql, {"session_id": session_id}).scalar()
        assert count == 0


def test_export_database_to_excel(tmp_path: Path) -> None:
    """Test exporting all database tables and records to an Excel workbook."""
    test_excel_file = tmp_path / "test_database_records.xlsx"
    saved_path = export_database_to_excel(output_path=test_excel_file)

    assert saved_path.exists()
    assert saved_path.stat().st_size > 0

    import openpyxl

    wb = openpyxl.load_workbook(str(saved_path))
    assert "Overview" in wb.sheetnames
    assert "alembic_version" in wb.sheetnames

    # Check alembic_version sheet contents
    ws_alembic = wb["alembic_version"]
    assert ws_alembic.cell(row=1, column=1).value == "version_num"
    assert ws_alembic.cell(row=2, column=1).value == "20260930_0003"


def run_standalone_inspection() -> None:
    """Print out database connection details, all table data retrieved via SQL, and export to Excel."""
    print("=" * 70)
    print("DATABASE CONTACT & ALL DATA RETRIEVAL REPORT")
    print("=" * 70)

    try:
        conn_info = check_db_connection()
        print(f"Connection Status : SUCCESS")
        print(f"Database Name     : {conn_info.get('db_name')}")
        print(f"Connected User    : {conn_info.get('user_name')}")
        print(f"PostgreSQL Version: {conn_info.get('pg_version')}")
    except Exception as exc:
        print(f"Connection Status : FAILED - {exc}")
        return

    print("-" * 70)
    print("Executing SQL queries to retrieve all database data...")
    print("-" * 70)

    all_data = get_all_database_data()
    for table_name, rows in all_data.items():
        print(f"\n[TABLE] {table_name} (Total rows: {len(rows)})")
        if not rows:
            print("  (Empty table - 0 records)")
        else:
            for idx, row in enumerate(rows, start=1):
                # Convert non-serializable objects (like datetimes/UUIDs) to string for display
                serializable_row = {k: str(v) if v is not None else None for k, v in row.items()}
                print(f"  Row {idx}: {json.dumps(serializable_row, indent=4)}")

    print("\n" + "-" * 70)
    print("Exporting all database records to Excel file in docs folder...")
    print("-" * 70)
    excel_path = export_database_to_excel()
    print(f"Saved Excel file to: {excel_path} ({excel_path.stat().st_size} bytes)")

    print("\n" + "=" * 70)
    print("Completed SQL queries for all tables in database successfully.")
    print("=" * 70)


if __name__ == "__main__":
    run_standalone_inspection()

